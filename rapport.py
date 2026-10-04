"""Rapport de corroboration exportable (Excel multi-onglets + CSV)."""
from __future__ import annotations

import io
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import __version__, retroaction
from .config import ANOMALIE, CHAMPS, CONFORME, HORS_PERIMETRE, JUSTIFIE
from .moteur import Resultat
from .normalisation import NORMALISATIONS
from .regles import REGLES, REGLES_ENREGISTREMENT
from .ia.local import HEURISTIQUES

COULEURS = {CONFORME: "D9EAD3", JUSTIFIE: "FFF2CC", ANOMALIE: "F4CCCC"}
ENTETE = PatternFill("solid", fgColor="1F3864")
BORD = Border(bottom=Side(style="thin", color="BFBFBF"))

COLS_ANOMALIES = ["ID", "Niveau priorité", "Priorité", "Matricule", "Affectation", "Champ cible", "Libellé",
                  "Valeur source", "Valeur attendue", "Valeur cible", "Verdict", "Confiance", "Origine de la décision",
                  "Règle", "Justification", "Diagnostic IA", "Action recommandée", "Preuves (signaux)",
                  "Trace du calcul", "Ligne source", "Ligne cible", "Accord LLM/local"]
COLS_JUSTIFIES = ["ID", "Matricule", "Affectation", "Champ cible", "Valeur source", "Valeur attendue", "Valeur cible",
                  "Confiance", "Origine de la décision", "Règle", "Justification", "Action recommandée",
                  "Preuves (signaux)", "Trace du calcul", "Ligne source", "Ligne cible"]
COLS_IA = ["ID", "Matricule", "Champ cible", "Valeur source", "Valeur attendue", "Valeur cible", "Verdict",
           "Confiance", "Origine de la décision", "Avis arbitre local", "Accord LLM/local", "Justification",
           "Preuves (signaux)", "Signature"]


def _feuilles(res: Resultat, corrections: list[dict] | None = None) -> dict[str, pd.DataFrame]:
    t = res.tableau()
    s = res.stats()
    anomalies = t[t["Verdict"] == ANOMALIE].sort_values(["Priorité", "Matricule"], ascending=[False, True])
    justifies = t[t["Verdict"] == JUSTIFIE].sort_values(["Champ cible", "Matricule"])
    ia = t[t["Origine de la décision"].str.startswith(("IA", "Règle apprise"))]

    synth = [("Rapport de corroboration CorroborIA", ""), ("Version", __version__),
             ("Exécuté le", res.horodatage), ("Moteur IA (niveau 3)", res.fournisseur),
             ("Interprétation R-DATE-DEBUT", res.params.mode_date_debut),
             ("Date de référence", str(res.params.date_reference)), ("", ""),
             ("Affectations source / cible", f"{s['affectations_source']} / {s['affectations_cible']}"),
             ("Contrôles réalisés (affectation × champ)", s["total"]),
             ("✔ Conformes", s[CONFORME]), ("≈ Écarts justifiés", s[JUSTIFIE]),
             ("✘ Anomalies à investiguer", s[ANOMALIE]),
             ("Employés touchés par une anomalie", s["employes_anomalie"]),
             ("Cas ambigus tranchés par l'IA", s["ambigus_ia"]), ("", "")]
    synth += [(f"Décisions – {k}", v) for k, v in s["origines"].items()]
    synth += [("", ""), ("Intégrité des fichiers sources (SHA-256, lecture seule)", "")]
    synth += [(k, v) for k, v in res.empreintes.items()]
    df_synth = pd.DataFrame(synth, columns=["Indicateur", "Valeur"])

    par_champ = (t.pivot_table(index="Champ cible", columns="Verdict", values="ID", aggfunc="count", fill_value=0)
                 .reindex(columns=[CONFORME, JUSTIFIE, ANOMALIE], fill_value=0).reset_index())

    regles = [{"Règle": r.id, "Titre": r.titre, "Description": r.description,
               "Champs cible": ", ".join(c.cible for c in CHAMPS if c.regle == r.id),
               "Référence mapping": ", ".join(sorted({x for c in CHAMPS if c.regle == r.id for x in c.cellules}))}
              for r in REGLES.values()]
    regles += [{"Règle": k, "Titre": "Complétude", "Description": v, "Champs cible": "(enregistrement)",
                "Référence mapping": "implicite"} for k, v in REGLES_ENREGISTREMENT.items()]
    regles += [{"Règle": k, "Titre": "Normalisation", "Description": v, "Champs cible": "tous",
                "Référence mapping": ""} for k, v in NORMALISATIONS.items()]
    regles += [{"Règle": k, "Titre": "Heuristique d'arbitrage IA locale", "Description": v, "Champs cible": "cas ambigus",
                "Référence mapping": ""} for k, v in HEURISTIQUES.items()]
    for r in retroaction.regles_apprises(corrections or []):
        regles.append({"Règle": r["id_regle"], "Titre": "Règle apprise (rétroaction)", "Description": r["texte_regle"],
                       "Champs cible": r["champ"], "Référence mapping": f"{r['auteur']} {r['horodatage']}"})
    regles += [{"Règle": "HORS PÉRIMÈTRE", "Titre": k, "Description": v, "Champs cible": "", "Référence mapping": ""}
               for k, v in HORS_PERIMETRE.items()]

    # index pandas i ↔ ligne Excel i + 2 (en-tête en ligne 1)
    app = pd.DataFrame([{"Matricule": p.matricule,
                         "Ligne source": None if p.i_src is None else p.i_src + 2,
                         "Ligne cible": None if p.i_cib is None else p.i_cib + 2,
                         "Score d'appariement": p.score, "Détail": p.details} for p in res.paires])
    return {
        "Synthèse": df_synth, "Par champ": par_champ,
        "Anomalies": anomalies[COLS_ANOMALIES], "Écarts justifiés": justifies[COLS_JUSTIFIES],
        "Analyse IA": ia[COLS_IA], "Détail complet": t,
        "Appariement": app, "Observations qualité": pd.DataFrame(res.observations),
        "Hypothèses": pd.DataFrame(res.validation_hypotheses),
        "Catalogue des règles": pd.DataFrame(regles),
        "Journal": pd.DataFrame({"Étape": res.journal}),
    }


def exporter_excel(res: Resultat, chemin: str | Path | None = None, corrections: list[dict] | None = None) -> bytes:
    buf = io.BytesIO()
    feuilles = _feuilles(res, corrections)
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for nom, df in feuilles.items():
            df.to_excel(xw, sheet_name=nom[:31], index=False)
            ws = xw.sheets[nom[:31]]
            for cell in ws[1]:
                cell.fill, cell.font = ENTETE, Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.freeze_panes = "B2"
            ws.auto_filter.ref = ws.dimensions
            for i, col in enumerate(df.columns, 1):
                larg = max([len(str(col))] + [min(len(str(v)), 60) for v in df[col].head(200)]) + 2
                ws.column_dimensions[get_column_letter(i)].width = min(max(larg, 9), 60)
            if "Verdict" in df.columns:
                iv = list(df.columns).index("Verdict") + 1
                for r in range(2, ws.max_row + 1):
                    v = ws.cell(r, iv).value
                    if v in COULEURS:
                        ws.cell(r, iv).fill = PatternFill("solid", fgColor=COULEURS[v])
            for r in ws.iter_rows(min_row=2):
                for c in r:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
                    c.border = BORD
        ws = xw.sheets["Synthèse"]
        ws["A1"].font = Font(bold=True, size=14)
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 55, 70
    data = buf.getvalue()
    if chemin:
        Path(chemin).parent.mkdir(parents=True, exist_ok=True)
        Path(chemin).write_bytes(data)
    return data


def exporter_csv(res: Resultat, chemin: str | Path | None = None) -> bytes:
    data = res.tableau().to_csv(index=False, sep=";").encode("utf-8-sig")
    if chemin:
        Path(chemin).parent.mkdir(parents=True, exist_ok=True)
        Path(chemin).write_bytes(data)
    return data
